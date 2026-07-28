"""
Build an editable Word (.docx) proposal draft from section text.

Uses python-docx. Output is bytes suitable for FastAPI StreamingResponse.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone


def build_proposal_docx(
    title: str,
    sections: dict[str, str],
    *,
    subtitle: str = "",
    footer_note: str = "Generated from scrubbed DFIR corpus — review before client use.",
) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title or "Proposal Draft", level=0)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        if p.runs:
            p.runs[0].font.size = Pt(11)

    meta = doc.add_paragraph(
        f"Generated at {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    if meta.runs:
        meta.runs[0].font.size = Pt(9)

    for heading, body in sections.items():
        if not body or not str(body).strip():
            continue
        label = heading.replace("_", " ").title()
        doc.add_heading(label, level=1)
        for para in str(body).strip().split("\n\n"):
            doc.add_paragraph(para.strip())

    doc.add_paragraph("")
    note = doc.add_paragraph(footer_note)
    if note.runs:
        note.runs[0].italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
